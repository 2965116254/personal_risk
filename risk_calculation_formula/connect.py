# -*- coding: utf-8 -*-
# @Project ：risk_v1 
# @FileName: connect.py
# @Author  : greenaut
# @Time    : 2026/2/28 9:37

import os
import re
import time
import logging
import sys
from croniter import croniter
from datetime import datetime
from minio import Minio
from minio.error import S3Error
import mysql.connector
from mysql.connector import Error
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from operator_capability_B import (
    calculate_worker_capacity_score,
    calculate_b3_score,
    calculate_b4_score
)
from operator_capability_C import calculate_work_risk_score
from operator_capability_D import calculate_device_risk_score
from name_cleaner import init_name_extractor, get_name_extractor
import config_loader

# ---------------------- 配置日志 ----------------------
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[logging.StreamHandler(sys.stdout)]
)
logger = logging.getLogger(__name__)

# ---------------------- CSV 文件路径（从配置获取） ----------------------
D_SCORE_PATH = config_loader.get_d_score_path()
FAULT_SCENE_PATH = config_loader.get_fault_scene_path()


def build_b1_criteria(principal_name):
    return f"{principal_name}：无违章记录"


def build_b2_criteria(member_text, name_extractor):
    if not member_text:
        return "无工作班成员信息"
    try:
        names = name_extractor.extract_names(member_text)
        if not names:
            return "无有效成员姓名"
        return "；".join([f"{name}：无违章记录" for name in names])
    except Exception as e:
        print(f"解析成员姓名失败: {e}")
        return "成员信息解析失败"


def merge_all_results(b_results, c_results, d_results, db_config=None, name_extractor=None):
    """
    合并 B、C、D 三类评分结果，计算 F = A+B+C+D（A 默认为 0）
    返回合并后的字典列表，每个字典包含所有评分字段及判断依据
    """
    if name_extractor is None:
        try:
            from name_cleaner import get_name_extractor
            name_extractor = get_name_extractor()
        except RuntimeError:
            print("错误：NameExtractor 未初始化，无法构建B2判断依据")
            name_extractor = None

    # 将各评分结果转为字典 {工作票id: 结果字典}
    b_dict = {item['工作票id']: item for item in b_results} if b_results else {}
    c_dict = {item['工作票id']: item for item in c_results} if c_results else {}
    d_dict = {item['工作票id']: item for item in d_results} if d_results else {}

    # 所有出现过的 ID 并集
    all_ids = set(b_dict.keys()) | set(c_dict.keys()) | set(d_dict.keys())

    # 批量查询缺失的 B 类信息（仅当 db_config 提供时）
    missing_ids = [tid for tid in all_ids if tid not in b_dict]
    missing_info = {}
    if missing_ids and db_config:
        try:
            connection = mysql.connector.connect(**db_config)
            cursor = connection.cursor(dictionary=True)
            placeholders = ','.join(['%s'] * len(missing_ids))
            query = f"""
                SELECT 
                    id,
                    work_principal_uname,
                    work_member_count,
                    whether_outer_dept,
                    work_member_uname
                FROM sp_pd_wticket_base
                WHERE id IN ({placeholders})
            """
            cursor.execute(query, missing_ids)
            for row in cursor.fetchall():
                missing_info[row['id']] = row
            print(f"批量补全了 {len(missing_info)} 张缺失B类评分的工作票信息")
        except Error as e:
            print(f"批量查询缺失工作票失败: {e}")
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    merged = []
    for tid in sorted(all_ids):  # 按 ID 排序输出
        item = {'工作票id': tid}

        # ---------- B 类数据 ----------
        if tid in b_dict:
            b = b_dict[tid]
            item.update({
                'B': b.get('B', 0),
                'B1': b.get('B1', 0),
                'B1判断依据（工作负责人名称）': b.get('B1判断依据（工作负责人名称）', ''),
                'B2': b.get('B2', 0),
                'B2判断依据（工作班组人员）': b.get('B2判断依据（工作班组人员）', ''),
                'B3': b.get('B3', 0),
                'B3作业总人数': b.get('B3作业总人数', 0),
                'B4': b.get('B4', 0),
                'B4人员性质（外来单位(1=是，2=否)）': b.get('B4人员性质（外来单位(1=是，2=否)）', '')
            })
        else:
            # 从 missing_info 或默认值构建
            info = missing_info.get(tid, {})
            principal = info.get('work_principal_uname') or ""
            member_count = info.get('work_member_count') or 0
            outer_dept = info.get('whether_outer_dept')
            member_text = info.get('work_member_uname') or ""
            b1_criteria = build_b1_criteria(principal)
            b2_criteria = build_b2_criteria(member_text, name_extractor)
            b3_score = calculate_b3_score(int(member_count)) if member_count else 0
            b4_score = calculate_b4_score(outer_dept)
            item.update({
                'B': 0, 'B1': 0, 'B1判断依据（工作负责人名称）': b1_criteria,
                'B2': 0, 'B2判断依据（工作班组人员）': b2_criteria,
                'B3': b3_score, 'B3作业总人数': member_count,
                'B4': b4_score, 'B4人员性质（外来单位(1=是，2=否)）': outer_dept
            })

        # ---------- C 类数据 ----------
        if tid in c_dict:
            c = c_dict[tid]
            item.update({
                'C': c.get('C', 0),
                'C1': c.get('C1', 0),
                'C1判断依据': c.get('C1判断依据', '缺气象数据'),
                'C2': c.get('C2', 0),
                'C2判断依据': c.get('C2判断依据', ''),
                'C3': c.get('C3', 0),
                'C3判断依据': c.get('C3判断依据', ''),
                'C4': c.get('C4', 0),
                'C4判断依据': c.get('C4判断依据', ''),
                'C5': c.get('C5', 0),
                'C5判断依据': c.get('C5判断依据', '')
            })
        else:
            item.update({
                'C': 0, 'C1': 0, 'C1判断依据': '无数据',
                'C2': 0, 'C2判断依据': '无数据',
                'C3': 0, 'C3判断依据': '无数据',
                'C4': 0, 'C4判断依据': '无数据',
                'C5': 0, 'C5判断依据': '无数据'
            })

        # ---------- D 类数据 ----------
        if tid in d_dict:
            d = d_dict[tid]
            item.update({
                'D值': d.get('D值', 0),
                '是否匹配典型场景': d.get('是否匹配典型场景', '否'),
                '事故后果描述': d.get('事故后果描述', ''),
                '判断依据': d.get('判断依据', '')
            })
        else:
            item.update({
                'D值': 0,
                '是否匹配典型场景': '否',
                '事故后果描述': '',
                '判断依据': '无D类评分数据'
            })

        # ---------- A 与 F ----------
        item['A'] = 0  # A 固定为 0
        item['F'] = item['B'] + item['C'] + item['D值']  # F = A+B+C+D, A=0

        merged.append(item)

    return merged


def save_merged_to_excel(merged_data, output_dir, filename):
    """保存合并结果到 Excel，返回完整文件路径"""
    if not merged_data:
        logger.warning("没有数据可保存")
        return None

    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
        logger.info(f"创建目录: {output_dir}")

    filepath = os.path.join(output_dir, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "作业票风险评分"

    # ----- 定义表头 -----
    headers = [
        '工作票id',
        'F（总风险值）',
        'A（默认0）',
        'B（作业人员能力风险值）',
        'B1',
        'B1判断依据（工作负责人名称）',
        'B2',
        'B2判断依据（工作班组人员）',
        'B3',
        'B3作业总人数',
        'B4',
        'B4人员性质（外来单位(1=是，2=否)）',
        'C（作业环境和时间影响风险值）',
        'C1',
        'C1判断依据',
        'C2',
        'C2判断依据',
        'C3',
        'C3判断依据',
        'C4',
        'C4判断依据',
        'C5',
        'C5判断依据',
        'D值',
        '是否匹配典型场景',
        '事故后果描述',
        '判断依据'
    ]

    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)

    # ----- 写入数据 -----
    for row_num, row_data in enumerate(merged_data, 2):
        col = 1
        ws.cell(row=row_num, column=col, value=row_data.get('工作票id', '')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('F', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('A', 0)); col += 1

        # B 类
        ws.cell(row=row_num, column=col, value=row_data.get('B', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('B1', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('B1判断依据（工作负责人名称）', '')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('B2', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('B2判断依据（工作班组人员）', '')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('B3', 0)); col += 1
        # B3作业总人数（文本格式）
        cell = ws.cell(row=row_num, column=col, value=str(row_data.get('B3作业总人数', 0)))
        cell.number_format = '@'; col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('B4', 0)); col += 1
        # B4人员性质（文本格式）
        cell = ws.cell(row=row_num, column=col, value=str(row_data.get('B4人员性质（外来单位(1=是，2=否)）', '')))
        cell.number_format = '@'; col += 1

        # C 类
        ws.cell(row=row_num, column=col, value=row_data.get('C', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C1', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C1判断依据', '')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C2', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C2判断依据', '')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C3', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C3判断依据', '')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C4', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C4判断依据', '')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C5', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('C5判断依据', '')); col += 1

        # D 类
        ws.cell(row=row_num, column=col, value=row_data.get('D值', 0)); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('是否匹配典型场景', '否')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('事故后果描述', '')); col += 1
        ws.cell(row=row_num, column=col, value=row_data.get('判断依据', '')); col += 1

    # ----- 调整列宽 -----
    col_widths = {
        1: 15, 2: 10, 3: 8,   # id, F, A
        4: 8, 5: 5, 6: 40, 7: 5, 8: 60, 9: 5, 10: 15, 11: 5, 12: 25,  # B类
        13: 8, 14: 5, 15: 5, 16: 5, 17: 40, 18: 5, 19: 40, 20: 5, 21: 40, 22: 5, 23: 25,  # C类
        24: 8, 25: 12, 26: 50, 27: 70   # D类
    }
    for col, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(filepath)
    logger.info(f"合并数据已保存到本地: {filepath}")
    logger.info(f"共保存 {len(merged_data)} 条记录")
    return filepath


def upload_to_minio(local_file_path, object_name=None):
    """上传文件到 MinIO，返回对象名称或 None"""
    cfg = config_loader.get_minio_config()
    try:
        client = Minio(
            cfg['endpoint'],
            access_key=cfg['access_key'],
            secret_key=cfg['secret_key'],
            secure=cfg.get('secure', True),
            region=cfg.get('region', None)
        )
        bucket = cfg['bucket_name']
        # 检查 bucket 是否存在，不存在则创建
        if not client.bucket_exists(bucket):
            client.make_bucket(bucket)
            logger.info(f"创建 bucket: {bucket}")

        # 确定对象名称
        if object_name is None:
            object_name = os.path.basename(local_file_path)
        # 如果配置了 object_prefix，则拼接前缀
        prefix = cfg.get('object_prefix', '').strip()
        if prefix:
            # 确保前缀末尾无多余斜杠，并拼接
            full_object_name = f"{prefix.rstrip('/')}/{object_name}"
        else:
            full_object_name = object_name

        # 上传
        client.fput_object(bucket, full_object_name, local_file_path)
        logger.info(f"文件已上传到 MinIO: {bucket}/{full_object_name}")
        return full_object_name
    except S3Error as e:
        logger.error(f"MinIO 上传失败: {e}")
        return None


def run_once():
    """执行一次评分计算、保存并上传"""
    logger.info("开始执行一次风险评估任务")
    db_config = config_loader.get_db_config()
    current_year = config_loader.get_current_year()
    limit = config_loader.get_query_limit()
    # output_dir = config_loader.get_output_dir()
    output_dir = os.path.join(os.getcwd(), "excel_output")
    # 初始化姓名提取器
    try:
        init_name_extractor(db_config)
    except Exception as e:
        logger.error(f"初始化姓名提取器失败: {e}")
        return

    # 检查 CSV 文件
    if not os.path.exists(D_SCORE_PATH):
        logger.error(f"D_score.csv 不存在: {D_SCORE_PATH}")
        return
    if not os.path.exists(FAULT_SCENE_PATH):
        logger.error(f"fault_scene.csv 不存在: {FAULT_SCENE_PATH}")
        return

    logger.info("开始计算 B 类评分...")
    b_results = calculate_worker_capacity_score(
        db_config, current_year, get_name_extractor(), limit=limit
    )

    logger.info("开始计算 C 类评分...")
    c_results = calculate_work_risk_score(db_config, limit=limit)

    logger.info("开始计算 D 类评分...")
    d_results = calculate_device_risk_score(
        db_config, D_SCORE_PATH, FAULT_SCENE_PATH, limit=limit
    )

    if not b_results and not c_results and not d_results:
        logger.warning("没有获取到任何数据，任务结束")
        return

    merged = merge_all_results(
        b_results, c_results, d_results,
        db_config=db_config,
        name_extractor=get_name_extractor()
    )

    logger.info(
        f"统计: B类 {len(b_results) if b_results else 0} 条, "
        f"C类 {len(c_results) if c_results else 0} 条, "
        f"D类 {len(d_results) if d_results else 0} 条, "
        f"合并后 {len(merged)} 条"
    )

    # 生成文件名
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f'作业人员风险评分_全类_{timestamp}.xlsx'
    local_file = save_merged_to_excel(merged, output_dir, filename)

    if local_file:
        # 上传到 MinIO
        upload_to_minio(local_file, object_name=filename)
        # 可选：删除本地文件（若不需要保留）
        os.remove(local_file)
    else:
        logger.error("本地文件未生成，无法上传")


def main():
    cfg = config_loader.get_config()
    schedule_cfg = cfg.get('schedule', {})

    # 优先使用 cron
    if 'cron' in schedule_cfg and schedule_cfg['cron']:
        cron_expr = schedule_cfg['cron']
        logger.info(f"使用 cron 表达式: {cron_expr}")
        base_time = datetime.now()
        cron = croniter(cron_expr, base_time)
        while True:
            next_time = cron.get_next(datetime)
            sleep_seconds = (next_time - datetime.now()).total_seconds()
            if sleep_seconds > 0:
                logger.info(f"等待 {sleep_seconds:.0f} 秒后执行下一次任务...")
                time.sleep(sleep_seconds)
            run_once()
    # 否则使用 interval
    elif 'interval' in schedule_cfg and schedule_cfg['interval']:
        interval = schedule_cfg['interval']
        logger.info(f"使用 interval 模式，每 {interval} 秒执行一次")
        while True:
            run_once()
            time.sleep(interval)
    else:
        logger.info("未配置定时，只执行一次")
        run_once()


if __name__ == "__main__":
    main()