# -*- coding: utf-8 -*-
"""
测试脚本：根据 test/data.json 中的作业计划编号查询工作任务，
通过异步大模型评估后生成问题描述 Excel（仅两列：作业计划编号、问题描述）。
规则使用 risk_calculation_formulav3.1/code 中的评估规则。
"""

import os
import sys
import json
import logging
from datetime import datetime

# ==================== 路径设置 ====================
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CODE_DIR = os.path.join(BASE_DIR, "code")
sys.path.insert(0, BASE_DIR)
sys.path.insert(0, CODE_DIR)

import config_loader
from risk_assessor import RiskAssessor
from excel_exporter import ExcelExporter
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, PatternFill, Font


def setup_logging():
    """配置日志输出到控制台，确保能看到大模型调用过程"""
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
        handlers=[logging.StreamHandler(sys.__stdout__)],
    )


def load_work_codes(json_path):
    """从 data.json 加载作业计划编号并去重（保持顺序）"""
    with open(json_path, 'r', encoding='utf-8') as f:
        codes = json.load(f)
    seen = set()
    unique_codes = []
    for c in codes:
        if c not in seen:
            seen.add(c)
            unique_codes.append(c)
    return unique_codes


def export_problem_description(results, output_path):
    """导出两列 Excel：作业计划编号、问题描述（复用 code 规则生成问题描述）"""
    exporter = ExcelExporter()

    wb = Workbook()
    ws = wb.active
    ws.title = '问题描述'

    # 表头
    ws.cell(row=1, column=1, value='作业计划编号')
    ws.cell(row=1, column=2, value='问题描述')

    for row_num, result in enumerate(results, 2):
        work_code = result.get('作业计划编号', '')
        detailed = result.get('详细评估结果', [])
        # 复用 ExcelExporter 中的规则生成问题描述（含标红富文本）
        judgment = exporter._generate_judgment(detailed)

        ws.cell(row=row_num, column=1, value=work_code)
        ws.cell(row=row_num, column=2, value=judgment)

        # 打印每条结果，确认模型在运行
        work_task = result.get('工作任务', '')
        print(f"[{row_num - 1}/{len(results)}] 作业计划编号: {work_code} | 工作任务: {work_task}")

    # ==================== 美化 ====================
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, name='微软雅黑')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 80

    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
            else:
                cell.alignment = left_align if cell.column == 2 else center_align

    ws.freeze_panes = 'A2'
    wb.save(output_path)
    print(f'\nExcel 已保存到: {output_path}')


def main():
    setup_logging()

    # 1. 加载作业计划编号
    json_path = os.path.join(BASE_DIR, 'test', 'data.json')
    work_codes = load_work_codes(json_path)
    print(f'从 data.json 加载了 {len(work_codes)} 个作业计划编号（去重后）')
    print(f'前5个编号: {work_codes[:5]}')

    # 2. 执行风险评估（内部通过 asyncio 异步并发调用大模型）
    db_config = config_loader.get_db_new_config()

    print('\n========== 开始风险评估（异步大模型调用）==========')
    assessor = RiskAssessor(db_config)
    results = assessor.assess(work_codes=work_codes, use_cache=False)

    if not results:
        print('没有获取到评估结果')
        return

    print(f'\n========== 风险评估完成，共 {len(results)} 条结果 ==========')

    # 3. 导出 Excel（仅作业计划编号 + 问题描述两列）
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_dir = os.path.join(BASE_DIR, 'output')
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, f'问题描述_{timestamp}.xlsx')

    print('\n========== 开始生成问题描述 Excel ==========')
    export_problem_description(results, output_path)


if __name__ == '__main__':
    main()
