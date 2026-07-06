# -*- coding: utf-8 -*-
"""
Excel 输出层
封装 Excel 文件的生成、数据写入和美化逻辑。
"""

import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from datetime import datetime
from typing import List, Optional, Dict

from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, PatternFill, Font
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from risk_rule_engine import RiskRuleEngine


class ExcelExporter:
    """Excel 导出类"""

    # 表头
    HEADERS = [
        '地市局',
        '作业计划编号',
        '工作内容',
        '典型基准风险值',
        '作业人员能力风险值',
        '作业环境和时间影响风险值',
        '电网、设备风险联动值',
        '总分',
        '风险等级',
        '问题描述',
        '违章代码',
        '违章条款',
    ]

    # 查询数据工作表表头
    QUERY_HEADERS = [
        '地市局',
        '作业计划编号',
        '作业人数',
        '负责人',
        '负责人违章',
        '班组成员',
        '班组成员违章',
        '监护人',
        '监护人违章',
        '工作内容',
        '工作任务',
    ]

    # 需要左对齐的列索引（1-based）
    LEFT_ALIGN_COLS = {5, 6, 10}

    def __init__(self, output_dir: str = None):
        self.output_dir = output_dir or os.path.join(os.getcwd(), 'output')
        self.engine = RiskRuleEngine()

    def export(self, results: List[Dict]) -> Optional[str]:
        """
        将评估结果导出到 Excel 文件
        :return: 文件路径或 None
        """
        if not results:
            print('没有数据，生成空 Excel（仅含表头）')

        os.makedirs(self.output_dir, exist_ok=True)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'数智问安-风险评估审核-{timestamp}.xlsx'
        filepath = os.path.join(self.output_dir, filename)

        # 预处理：在内存中过滤无违章数据，避免逐行 delete_rows（大数据量时极慢）
        filtered_results = []
        for result in results:
            detailed = result.get('详细评估结果', [])
            has_red = any(
                d.get('风险值得分', 0) > d.get('客户填入分值', 0)
                for d in detailed
            )
            if has_red:
                filtered_results.append(result)

        wb = Workbook()
        ws = wb.active
        ws.title = '风险评估结果'

        # 写入表头
        for col, header in enumerate(self.HEADERS, 1):
            ws.cell(row=1, column=col, value=header)

        # 写入数据（已过滤，无需后续删除操作）
        for row_num, result in enumerate(filtered_results, 2):
            self._write_row(ws, row_num, result)

        # 美化（合并遍历，含列宽采样优化）
        self._beautify(ws, len(filtered_results))

        # 写入查询数据工作表（使用全量数据，不过滤）
        self._write_query_data_sheet(wb, results)

        # 保存
        wb.save(filepath)

        print(f'数据已保存到: {filepath}')
        print(f'共保存了 {len(filtered_results)} 条工作计划编号的评估结果（已过滤无违章数据）')
        return filepath

    def _write_row(self, ws, row_num: int, result: Dict):
        """写入单行数据"""
        detailed = result.get('详细评估结果', [])

        # 构建评估因子查找表（避免每次 _get_detail_score/_get_detail_customer 都线性遍历）
        detail_map = {}
        for d in detailed:
            detail_map[d.get('评估因子', '')] = d

        def _ds(factor: str) -> float:
            d = detail_map.get(factor)
            return float(d.get('风险值得分', 0)) if d else 0.0

        def _dc(factor: str) -> float:
            d = detail_map.get(factor)
            return float(d.get('客户填入分值', 0)) if d else 0.0

        # 1. 地市局
        ws.cell(row=row_num, column=1, value=result.get('地市局', ''))

        # 2. 作业计划编号
        ws.cell(row=row_num, column=2, value=result.get('作业计划编号', ''))

        # 3. 工作内容
        ws.cell(row=row_num, column=3, value=result.get('工作内容', ''))

        # 4. 典型基准风险值（A值）
        benchmark_data = result.get('基准关系', [])
        a_value = result.get('A（基准风险值）', 0)
        if benchmark_data:
            benchmark_items = '\n'.join(
                f"{b.get('benchmark_name', '')}: {b.get('risk_value', 0)}分"
                for b in benchmark_data if b.get('benchmark_name')
            )
            benchmark_text = f"典型基准风险值 = {round(a_value)}分\n{benchmark_items}"
        else:
            benchmark_text = '无基准项目'
        ws.cell(row=row_num, column=4, value=benchmark_text)

        # 5. 作业人员能力风险值 (B)
        b_score = result.get('B（作业人员能力风险值）', 0)
        pg_model = _ds('现场作业负责人（含小组工作负责人）及监护人（含专职监护人）安全意识')
        pg_customer = _dc('现场作业负责人（含小组工作负责人）及监护人（含专职监护人）安全意识')
        mb_model = _ds('主要工作班成员(辅助工除外)安全意识')
        mb_customer = _dc('主要工作班成员(辅助工除外)安全意识')
        cnt_model = _ds('作业总人数')
        cnt_customer = _dc('作业总人数')
        nat_model = _ds('负责人的人员性质')
        nat_customer = _dc('负责人的人员性质')

        customer_b = int(pg_customer) + int(mb_customer) + int(cnt_customer) + int(nat_customer)
        b_segments = [
            (f"总分: 【模型评估】{b_score}分、【人工评估】{customer_b}分\n", b_score > customer_b),
            (f"现场作业负责人（含小组工作负责人）及监护人（含专职监护人）安全意识: 【模型评估】{pg_model}分、【人工评估】{int(pg_customer)}分\n", int(pg_model) > int(pg_customer)),
            (f"主要工作班成员(辅助工除外)安全意识: 【模型评估】{mb_model}分、【人工评估】{int(mb_customer)}分\n", int(mb_model) > int(mb_customer)),
            (f"作业总人数: 【模型评估】{cnt_model}分、【人工评估】{int(cnt_customer)}分\n", int(cnt_model) > int(cnt_customer)),
            (f"负责人的人员性质: 【模型评估】{nat_model}分、【人工评估】{int(nat_customer)}分", int(nat_model) > int(nat_customer)),
        ]
        ws.cell(row=row_num, column=5, value=self._make_rich_text_opt(b_segments))

        # 6. 作业环境和时间影响风险值 (C)
        c_score = result.get('C（作业环境和时间影响风险值）', 0)
        loc_model = _ds('作业地段')
        loc_customer = _dc('作业地段')
        typ_model = _ds('作业类型')
        typ_customer = _dc('作业类型')
        wea_model = _ds('天气')
        wea_customer = _dc('天气')
        tp_model = _ds('作业时段')
        tp_customer = _dc('作业时段')

        customer_c = int(loc_customer) + int(typ_customer) + int(wea_customer) + int(tp_customer)
        c_segments = [
            (f"总分: 【模型评估】{c_score}分、【人工评估】{customer_c}分\n", c_score > customer_c),
            (f"作业地段: 【模型评估】{loc_model}分、【人工评估】{int(loc_customer)}分\n", int(loc_model) > int(loc_customer)),
            (f"作业类型: 【模型评估】{typ_model}分、【人工评估】{int(typ_customer)}分\n", int(typ_model) > int(typ_customer)),
            (f"天气: 【模型评估】{wea_model}分、【人工评估】{int(wea_customer)}分\n", int(wea_model) > int(wea_customer)),
            (f"作业时段: 【模型评估】{tp_model}分、【人工评估】{int(tp_customer)}分", int(tp_model) > int(tp_customer)),
        ]
        ws.cell(row=row_num, column=6, value=self._make_rich_text_opt(c_segments))

        # 7. 电网、设备风险联动值 (D)
        d_score = result.get('D（电网、设备风险联动值）', 0)
        d_items = result.get('D_items', [])
        if d_items:
            d_lines = [f"电网、设备风险联动值 = 总分（{int(d_score)}分）"]
            for item in d_items:
                d_lines.append(f"{item.get('评估因子', '')}：{int(item.get('风险值得分', 0))}分")
            ws.cell(row=row_num, column=7, value='\n'.join(d_lines))
        else:
            ws.cell(row=row_num, column=7, value=f"总分：{int(d_score)}")

        # 8. 总分（含基准风险值A）
        a_value = result.get('A（基准风险值）', 0)
        f_score = result.get('F（总风险值）', 0)
        total_customer = int(customer_b + customer_c + a_value + d_score)
        ws.cell(row=row_num, column=8, value=self._make_rich_text_opt([
            (f"【模型评估】{int(f_score)}分 【人工评估】{total_customer}分", int(f_score) > total_customer),
        ]))

        # 9. 风险等级
        model_level = self.engine.get_risk_level(f_score)
        customer_level = self.engine.get_risk_level(total_customer)
        ws.cell(row=row_num, column=9, value=f"【模型评估】{model_level} 【人工评估】{customer_level}")

        # 10. 问题描述
        ws.cell(row=row_num, column=10, value=self._generate_judgment(detailed))

        # 11. 违章代码
        ws.cell(row=row_num, column=11, value='D10')

        # 12. 违章条款
        ws.cell(row=row_num, column=12, value='信息系统的作业信息填报不正确、不规范')

    def _generate_judgment(self, detailed: List[Dict]):
        """生成规则判断结果，返回 CellRichText（模型评估>人工评估的片段标红）"""
        differences = [d for d in detailed if d.get('风险值得分', 0) != d.get('客户填入分值', 0)]
        if not differences:
            return '模型评估结果与客户填写结果一致'  # 不需标红时用纯文本，避免 RichText 开销

        segments = []
        for diff in differences:
            factor = diff.get('评估因子', '')
            rule_score = diff.get('风险值得分', 0)
            customer_score = diff.get('客户填入分值', 0)
            evaluation_result = diff.get('评估结果', '')
            llm_rule = diff.get('大模型命中规则', '')
            llm_keywords = diff.get('大模型命中关键词', '')
            llm_inferred = diff.get('大模型推断依据', '')

            simplified = self._simplify_factor(factor)
            is_red = rule_score > customer_score
            segments.append((f"{simplified}: 模型评估{rule_score}分，人工评估{customer_score}分\n", is_red))

            if factor in ('作业地段', '作业类型'):
                reason = f"规则判断依据：模型判断{simplified}为{rule_score}分"
                if llm_rule:
                    reason += f"，{llm_rule}"
                if llm_keywords:
                    reason += f"，工作内容中命中关键词：{llm_keywords}"
                if llm_inferred:
                    reason += f"（{llm_inferred}）"
                segments.append((f"{reason}\n", False))
            else:
                if self._is_awareness_factor(factor):
                    if evaluation_result and evaluation_result != '无人员':
                        segments.append((f"规则判断依据：{evaluation_result}\n", False))
                        segments.append((f"评分说明：根据安全意识评分规则——有A类违章得6分、有B类违章得5分、有C类违章得3分、有D类违章3次及以上得2分、D类违章不足3次或无违章得0分，取最高分作为该项得分\n", False))
                    else:
                        segments.append((f"规则判断依据：无相关作业人员信息，根据规则默认得{rule_score}分\n", False))
                elif factor == '作业总人数':
                    count_str = evaluation_result if evaluation_result and evaluation_result != '未知' else '0'
                    result_line = f"规则判断依据：作业总人数为{count_str}人"
                    try:
                        count = int(count_str)
                        thresholds = [
                            (50, 15, '≥50人'),
                            (24, 8, '24-49人'),
                            (16, 5, '16-23人'),
                            (8, 3, '8-15人'),
                            (5, 1, '5-7人'),
                        ]
                        rule_desc = '，根据人数评分规则：'
                        for threshold, score, desc in thresholds:
                            if count >= threshold:
                                rule_desc += f'{desc}得{score}分'
                                break
                        else:
                            rule_desc += '不足5人得0分'
                        result_line += f"{rule_desc}"
                    except (ValueError, TypeError):
                        result_line += '，无法根据人数确定分值，默认得0分'
                    segments.append((f"{result_line}\n", False))
                elif factor == '负责人的人员性质':
                    nature_explanation = {
                        '本单位-系统内人员': '得0分',
                        '总包单位作业': '得3分',
                        '外单位-系统外人员': '得5分',
                        '未知人员性质': '默认得0分',
                    }
                    explanation = nature_explanation.get(evaluation_result, f'对应规则得{rule_score}分')
                    segments.append((f'规则判断依据：作业主体的人员性质为"{evaluation_result}"，根据人员性质评分规则——{explanation}\n', False))
                elif factor == '作业时段':
                    segments.append((f'规则判断依据：作业时段被判定为"{evaluation_result}"，根据作业时段评分规则得{rule_score}分\n', False))
                elif evaluation_result:
                    segments.append((f"规则判断依据：{evaluation_result}\n", False))
                else:
                    segments.append((f"规则判断依据：根据{simplified}评估规则，计算得分为{rule_score}分\n", False))

        # 去掉最后一个换行符（如果有）
        if segments and segments[-1][0].endswith('\n'):
            last_text, last_red = segments[-1]
            segments[-1] = (last_text.rstrip('\n'), last_red)

        return self._make_rich_text_opt(segments)

    @staticmethod
    def _is_awareness_factor(factor: str) -> bool:
        """判断是否为安全意识评估因子"""
        return factor in (
            '现场作业负责人（含小组工作负责人）及监护人（含专职监护人）安全意识',
            '主要工作班成员(辅助工除外)安全意识',
        )

    @staticmethod
    def _simplify_factor(factor: str) -> str:
        """简化评估因子名称"""
        if '现场作业负责人' in factor:
            return '现场作业负责人及监护人安全意识'
        if '主要工作班成员' in factor:
            return '主要工作班成员安全意识'
        if '负责人的人员性质' in factor:
            return '人员性质'
        return factor

    @staticmethod
    def _get_detail_score(detailed: List[Dict], factor: str) -> float:
        for d in detailed:
            if d.get('评估因子', '') == factor:
                return float(d.get('风险值得分', 0))
        return 0.0

    @staticmethod
    def _get_detail_customer(detailed: List[Dict], factor: str) -> float:
        for d in detailed:
            if d.get('评估因子', '') == factor:
                return float(d.get('客户填入分值', 0))
        return 0.0

    # ==================== 富文本工具 ====================

    @staticmethod
    def _make_rich_text(segments):
        """
        根据 (text, is_red) 列表构建 CellRichText 对象
        segments: List[Tuple[str, bool]]  每个元素为 (文本, 是否标红)
        """
        red_font = InlineFont(rFont='微软雅黑', sz=10, color='FF0000')
        normal_font = InlineFont(rFont='微软雅黑', sz=10, color='000000')
        blocks = []
        for text, is_red in segments:
            font = red_font if is_red else normal_font
            blocks.append(TextBlock(font, text))
        return CellRichText(*blocks)

    @staticmethod
    def _make_rich_text_opt(segments):
        """
        优化版富文本构建：当没有任何片段需要标红时，直接返回纯文本字符串，
        避免创建 CellRichText 对象（大幅减少 openpyxl 内部开销）。
        segments: List[Tuple[str, bool]]  每个元素为 (文本, 是否标红)
        """
        has_red = any(is_red for _, is_red in segments)
        if not has_red:
            return ''.join(text for text, _ in segments)
        return ExcelExporter._make_rich_text(segments)

    # ==================== Excel 美化 ====================

    def _beautify(self, ws, row_count: int):
        """
        美化 Excel 工作表（在保存前执行，保留富文本格式）

        优化：
        - 单次遍历完成样式设置 + 列宽计算
        - 大数据量时列宽采用采样策略（避免全量 encode 计算）
        """
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, name='微软雅黑')
        content_font = Font(name='微软雅黑', size=10)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

        # 列宽采样策略：超过阈值时只采样部分行以节省时间
        SAMPLE_THRESHOLD = 200
        SAMPLE_HEAD = 100  # 前 N 行始终采样
        # 每隔 N 行采样一行
        sample_interval = max(1, row_count // 100) if row_count > SAMPLE_THRESHOLD else 1

        col_widths = [0] * (len(self.HEADERS) + 1)  # 1-based 索引

        for row_idx, row in enumerate(ws.iter_rows()):
            row_num = row_idx + 1

            # 判断是否参与列宽计算（采样逻辑）
            should_sample = (
                row_count <= SAMPLE_THRESHOLD
                or row_num == 1  # 表头始终参与
                or row_num <= SAMPLE_HEAD
                or (row_num - SAMPLE_HEAD) % sample_interval == 0
            )

            for cell in row:
                # ── 设置样式 ──
                cell.border = thin_border
                if row_num == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align if cell.column in self.LEFT_ALIGN_COLS else center_align
                    if not isinstance(cell.value, CellRichText):
                        cell.font = content_font

                # ── 列宽计算（采样） ──
                if should_sample and cell.value:
                    try:
                        length = len(str(cell.value).encode('gbk'))
                        if length > col_widths[cell.column]:
                            col_widths[cell.column] = length
                    except Exception:
                        pass

        # 设置列宽
        for col_idx, width in enumerate(col_widths[1:], 1):
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[col_letter].width = min(max(width + 2, 12), 50)

        ws.freeze_panes = 'A2'

    # ==================== 查询数据工作表 ====================

    def _write_query_data_sheet(self, wb, results):
        """写入查询数据工作表"""
        ws = wb.create_sheet('查询数据')

        # 写入表头
        for col, header in enumerate(self.QUERY_HEADERS, 1):
            ws.cell(row=1, column=col, value=header)

        # 写入数据
        row_num = 2
        for result in results:
            query_data = result.get('查询数据', {})
            if not query_data:
                continue
            for col, header in enumerate(self.QUERY_HEADERS, 1):
                value = query_data.get(header, '')
                if value is None:
                    value = ''
                ws.cell(row=row_num, column=col, value=value)
            row_num += 1

        # 美化
        self._beautify_query_sheet(ws, row_num - 1)

    def _beautify_query_sheet(self, ws, row_count):
        """美化查询数据工作表"""
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, name='微软雅黑')
        content_font = Font(name='微软雅黑', size=10)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
        # 需要左对齐的列（工作内容、工作任务等文本列）
        query_left_align_cols = {6, 7, 10, 11}

        for row in ws.iter_rows():
            for cell in row:
                cell.border = thin_border
                if cell.row == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align
                else:
                    cell.font = content_font
                    cell.alignment = left_align if cell.column in query_left_align_cols else center_align

        # 设置列宽
        col_widths = [18, 22, 10, 12, 14, 18, 18, 12, 14, 30, 30]
        for col_idx, width in enumerate(col_widths, 1):
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = 'A2'